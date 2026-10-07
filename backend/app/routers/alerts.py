from __future__ import annotations

from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..demo_scope import is_demo_post
from ..display_filters import sanitize_source_label, should_hide_post
from ..models import AlertRecord, AnalysisResult, ArchiveRecord, Post
from ..schemas import ReviewRequest


router = APIRouter(prefix="/api/alerts", tags=["alerts"])

RISK_LEVEL_DISPLAY = {
    "low": "低风险",
    "medium": "中风险",
    "high": "高风险",
}

ALERT_STATUS_DISPLAY = {
    "pending": "待处理",
    "reviewed": "已复核",
    "archived": "已归档",
}


def _normalize_alert_status(status: str | None) -> str:
    value = str(status or "").strip().lower()
    return value if value in ALERT_STATUS_DISPLAY else "pending"


def _resolve_analysis_and_post(db: Session, alert: AlertRecord):
    analysis = db.query(AnalysisResult).filter(AnalysisResult.id == alert.analysis_id).first()
    if not analysis:
        return None, None
    post = db.query(Post).filter(Post.id == analysis.post_id_fk).first()
    if not post or not is_demo_post(post) or should_hide_post(post):
        return analysis, None
    return analysis, post


def _format_datetime(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _collect_alert_rows(db: Session):
    alerts = db.query(AlertRecord).order_by(AlertRecord.created_at.desc()).all()
    result = []
    for alert in alerts:
        analysis, post = _resolve_analysis_and_post(db, alert)
        if not analysis or not post:
            continue

        alert_status = _normalize_alert_status(alert.alert_status)
        risk_level_key = (analysis.risk_level or "low").lower()
        result.append(
            {
                "alert_id": alert.id,
                "alert_status": ALERT_STATUS_DISPLAY.get(alert_status, "待处理"),
                "alert_status_key": alert_status,
                "review_note": alert.review_note,
                "created_at": alert.created_at,
                "created_at_display": _format_datetime(alert.created_at),
                "post_id": post.post_id,
                "title": post.title,
                "excerpt": post.excerpt,
                "source": sanitize_source_label(post.source, post.forum_name),
                "content": post.content,
                "post_url": post.post_url,
                "risk_level": RISK_LEVEL_DISPLAY.get(risk_level_key, "低风险"),
                "risk_level_key": risk_level_key,
                "labels": analysis.labels,
                "risk_score": analysis.risk_score,
                "confidence": analysis.confidence,
                "matched_keywords": analysis.matched_keywords,
            }
        )

    risk_rank = {"high": 0, "medium": 1, "low": 2}
    result.sort(key=lambda item: (risk_rank.get(item["risk_level_key"], 9), -(item["risk_score"] or 0)))
    return result


@router.get("")
def get_alerts(db: Session = Depends(get_db)):
    return _collect_alert_rows(db)


@router.get("/export")
def export_alerts(db: Session = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="Excel export dependency missing. Install openpyxl from backend requirements.",
        ) from exc

    rows = _collect_alert_rows(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Risk Alert List"

    headers = [
        "预警ID",
        "帖子ID",
        "帖子标题",
        "摘要",
        "来源",
        "风险等级",
        "标签",
        "风险分",
        "置信度",
        "命中证据",
        "状态",
        "处理备注",
        "生成时间",
        "原帖链接",
    ]
    sheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="FFE7BF")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(
            [
                row["alert_id"],
                row["post_id"],
                row["title"],
                row["excerpt"] or row["content"],
                row["source"],
                row["risk_level"],
                row["labels"],
                row["risk_score"],
                row["confidence"],
                row["matched_keywords"],
                row["alert_status"],
                row["review_note"],
                row["created_at_display"],
                row["post_url"],
            ]
        )

    column_widths = {
        "A": 10,
        "B": 14,
        "C": 28,
        "D": 42,
        "E": 18,
        "F": 12,
        "G": 24,
        "H": 10,
        "I": 10,
        "J": 28,
        "K": 14,
        "L": 24,
        "M": 20,
        "N": 28,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"risk_alert_list_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("/{alert_id}/review")
def review_alert(alert_id: int, data: ReviewRequest, db: Session = Depends(get_db)):
    alert = db.query(AlertRecord).filter(AlertRecord.id == alert_id).first()
    analysis, post = _resolve_analysis_and_post(db, alert) if alert else (None, None)
    if alert and (not analysis or not post):
        raise HTTPException(status_code=404, detail="alert not found")
    if not alert:
        raise HTTPException(status_code=404, detail="alert not found")
    alert.alert_status = _normalize_alert_status(data.alert_status)
    alert.review_note = data.review_note
    db.commit()
    return {"message": "处理成功"}


@router.post("/{alert_id}/archive")
def archive_alert(alert_id: int, db: Session = Depends(get_db)):
    alert = db.query(AlertRecord).filter(AlertRecord.id == alert_id).first()
    analysis, post = _resolve_analysis_and_post(db, alert) if alert else (None, None)
    if alert and (not analysis or not post):
        raise HTTPException(status_code=404, detail="alert not found")
    if not alert:
        raise HTTPException(status_code=404, detail="alert not found")
    archive = ArchiveRecord(alert_id=alert_id, archive_reason="已处理完成")
    alert.alert_status = "archived"
    db.add(archive)
    db.commit()
    return {"message": "归档成功"}
