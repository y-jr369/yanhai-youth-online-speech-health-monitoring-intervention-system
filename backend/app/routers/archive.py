from __future__ import annotations

from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..demo_scope import is_demo_post
from ..display_filters import sanitize_forum_name, sanitize_source_label, should_hide_post
from ..models import AlertRecord, AnalysisResult, ArchiveRecord, InterventionRecord, Post


router = APIRouter(prefix="/api/archive", tags=["archive"])


def _format_datetime(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _collect_archive_rows(db: Session):
    rows = db.query(ArchiveRecord).order_by(ArchiveRecord.archived_at.desc()).all()
    result = []

    for archive in rows:
        alert = db.query(AlertRecord).filter(AlertRecord.id == archive.alert_id).first()
        if not alert:
            continue

        analysis = db.query(AnalysisResult).filter(AnalysisResult.id == alert.analysis_id).first()
        if not analysis:
            continue

        post = db.query(Post).filter(Post.id == analysis.post_id_fk).first()
        if not post or not is_demo_post(post) or should_hide_post(post):
            continue

        intervention = (
            db.query(InterventionRecord)
            .filter(InterventionRecord.alert_id == archive.alert_id)
            .order_by(InterventionRecord.id.desc())
            .first()
        )

        result.append(
            {
                "id": archive.id,
                "alert_id": archive.alert_id,
                "archive_reason": archive.archive_reason,
                "archived_by": archive.archived_by,
                "archived_at": archive.archived_at,
                "archived_at_display": _format_datetime(archive.archived_at),
                "risk_level": analysis.risk_level,
                "labels": analysis.labels,
                "risk_score": analysis.risk_score,
                "confidence": analysis.confidence,
                "post_id": post.post_id,
                "title": post.title,
                "excerpt": post.excerpt,
                "author_name": post.author_name,
                "source": sanitize_source_label(post.source, post.forum_name),
                "forum_name": sanitize_forum_name(post.forum_name),
                "post_url": post.post_url,
                "publish_time": _format_datetime(post.publish_time),
                "campus_process_no": intervention.campus_process_no if intervention else "",
                "handler": intervention.handler if intervention else "",
                "process_status": intervention.process_status if intervention else "",
                "follow_up_note": intervention.follow_up_note if intervention else "",
                "completed_at": _format_datetime(intervention.completed_at) if intervention else "",
            }
        )

    return result


@router.get("")
def get_archives(db: Session = Depends(get_db)):
    return _collect_archive_rows(db)


@router.get("/export")
def export_archives(db: Session = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="Excel export dependency missing. Install openpyxl from backend requirements.",
        ) from exc

    rows = _collect_archive_rows(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Archive Records"

    headers = [
        "归档ID",
        "预警ID",
        "风险等级",
        "标签",
        "风险分",
        "置信度",
        "帖子ID",
        "帖子标题",
        "摘要",
        "作者",
        "来源",
        "贴吧",
        "原帖链接",
        "发布时间",
        "归档原因",
        "归档人",
        "归档时间",
        "工单号",
        "处理人",
        "处理状态",
        "跟进备注",
        "完成时间",
    ]
    sheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="DCEBFF")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(
            [
                row["id"],
                row["alert_id"],
                row["risk_level"],
                row["labels"],
                row["risk_score"],
                row["confidence"],
                row["post_id"],
                row["title"],
                row["excerpt"],
                row["author_name"],
                row["source"],
                row["forum_name"],
                row["post_url"],
                row["publish_time"],
                row["archive_reason"],
                row["archived_by"],
                row["archived_at_display"],
                row["campus_process_no"],
                row["handler"],
                row["process_status"],
                row["follow_up_note"],
                row["completed_at"],
            ]
        )

    column_widths = {
        "A": 10,
        "B": 10,
        "C": 12,
        "D": 24,
        "E": 10,
        "F": 10,
        "G": 16,
        "H": 28,
        "I": 40,
        "J": 14,
        "K": 18,
        "L": 18,
        "M": 28,
        "N": 20,
        "O": 20,
        "P": 14,
        "Q": 20,
        "R": 18,
        "S": 14,
        "T": 14,
        "U": 28,
        "V": 20,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"archive_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Access-Control-Expose-Headers": "Content-Disposition",
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
